import json
import os
import re
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from pypdf import PdfReader

st.set_page_config(page_title="Resume ATS Checker", page_icon="📄", layout="wide")
st.title("📄 AI Resume ATS Checker")
st.write(
    "Upload your resume and optionally paste a job description to get an estimated ATS-style score, "
    "keyword-gap analysis, and practical improvement suggestions."
)
st.caption(
    "This is an AI estimate, not a score from a real applicant-tracking system or a hiring decision."
)


def extract_text(uploaded_file) -> str:
    filename = uploaded_file.name.lower()
    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif filename.endswith(".docx"):
        document = Document(uploaded_file)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif filename.endswith(".txt"):
        text = uploaded_file.getvalue().decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(
            "No readable text was found. Scanned PDFs may need OCR before you upload them."
        )
    return text


def get_api_key() -> str | None:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.getenv("GEMINI_API_KEY")


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict[str, Any]:
    client = genai.Client(api_key=api_key)
    prompt = "\n".join(
        [
            "You are a careful resume reviewer. Analyze the resume against the job description, if provided.",
            "Return ONLY valid JSON with these keys:",
            "ats_score: integer from 0 to 100",
            "score_explanation: short string explaining this is an estimate",
            "strengths: array of 3 to 6 concise strings",
            "improvements: array of 4 to 8 actionable strings",
            "missing_keywords: array of relevant keywords or skills not clearly evidenced in the resume",
            "formatting_checks: array of observations about ATS readability, sections, dates, and layout",
            "summary: concise overall assessment",
            "rewritten_bullet_examples: array of up to 3 improved resume bullet examples; do not invent achievements or numbers",
            "Rules:",
            "If no job description is provided, assess general resume clarity and ATS readability and say job-specific matching cannot be assessed.",
            "Do not claim to know the score from any real ATS product.",
            "Do not infer protected traits or judge the applicant based on them.",
            "Never invent qualifications, employers, credentials, metrics, or accomplishments. Mark examples needing user verification.",
            "Keep feedback respectful, specific, and practical.",
            "RESUME TEXT:",
            resume_text[:24000],
            "JOB DESCRIPTION:",
            job_description[:12000] if job_description.strip() else "Not provided.",
        ]
    )
    response = client.models.generate_content(
        model="gemini-2.5-flash",
        contents=prompt,
        config={"response_mime_type": "application/json"},
    )
    raw = (response.text or "").strip()
    if not raw:
        raise RuntimeError("Gemini returned an empty response. Please try again.")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if not match:
            raise RuntimeError("The AI response was not valid JSON. Please try again.")
        data = json.loads(match.group(0))
